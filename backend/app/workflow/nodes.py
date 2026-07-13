from pathlib import Path

from docx import Document
from docx.shared import Pt

from ..parsers.document_parser import parse_document
from ..services.brd_generator import build_section, chunk_text, generate_section_content
from .sections import BRD_SECTIONS
from .state import WorkflowState


def parse_document_node(state: WorkflowState) -> WorkflowState:
    file_path = Path(state["file_path"])
    try:
        raw_text = parse_document(file_path)
        return {
            "raw_text": raw_text,
            "status": "parsed",
            "error": None,
        }
    except Exception as exc:  # noqa: BLE001
        return {"status": "failed", "error": str(exc)}


def chunk_document_node(state: WorkflowState) -> WorkflowState:
    chunks = chunk_text(state.get("raw_text", ""))
    return {
        "chunks": chunks,
        "section_titles": BRD_SECTIONS,
        "current_section_index": 0,
        "sections": [],
        "status": "chunked",
    }


def generate_section_node(state: WorkflowState) -> WorkflowState:
    index = state.get("current_section_index", 0)
    titles = state.get("section_titles", BRD_SECTIONS)
    if index >= len(titles):
        return {"status": "sections_complete"}

    title = titles[index]
    feedback = state.get("approval_feedback")
    section_id = title.lower().replace(" ", "-")
    existing = next((section for section in state.get("sections", []) if section["id"] == section_id), None)
    approved_sections = [
        section for section in state.get("sections", []) if section.get("status") == "approved"
    ]
    content = generate_section_content(
        title=title,
        chunks=state.get("chunks", []),
        feedback=feedback,
        approved_sections=approved_sections,
        filename=state.get("filename"),
        previous_draft=existing["content"] if existing else None,
    )
    section = build_section(title, content, status="awaiting_approval")
    return {
        "sections": [section],
        "status": "awaiting_approval",
        "approval_decision": None,
        "approval_feedback": None,
    }


def await_approval_node(state: WorkflowState) -> WorkflowState:
    return {"status": "awaiting_approval"}


def process_approval_node(state: WorkflowState) -> WorkflowState:
    decision = state.get("approval_decision")
    index = state.get("current_section_index", 0)
    titles = state.get("section_titles", BRD_SECTIONS)
    title = titles[index]
    section_id = title.lower().replace(" ", "-")

    sections = list(state.get("sections", []))
    current = next((section for section in sections if section["id"] == section_id), None)
    if current is None:
        return {"status": "failed", "error": "Current section not found"}

    if decision == "reject":
        current["status"] = "rejected"
        current["version"] = current.get("version", 1) + 1
        return {
            "sections": [current],
            "status": "regenerate",
            "approval_decision": None,
        }

    current["status"] = "approved"
    next_index = index + 1
    updates: WorkflowState = {
        "sections": [current],
        "current_section_index": next_index,
        "approval_decision": None,
        "approval_feedback": None,
    }
    if next_index >= len(titles):
        updates["status"] = "sections_complete"
    else:
        updates["status"] = "generating"
    return updates


def compile_brd_node(state: WorkflowState) -> WorkflowState:
    workflow_id = state["workflow_id"]
    output_dir = Path("data/outputs")
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{workflow_id}.docx"

    document = Document()
    document.add_heading("BUSINESS REQUIREMENTS DOCUMENT", level=0)
    document.add_paragraph(f"Generated from: {state.get('filename', 'uploaded document')}")
    document.add_paragraph("Status: Approved")

    for section in state.get("sections", []):
        document.add_heading(section["title"], level=1)
        for line in section["content"].splitlines():
            if line.startswith("### "):
                document.add_heading(line.replace("### ", ""), level=2)
            elif line.startswith("|") or line.startswith("---"):
                continue
            elif line.startswith("- **"):
                paragraph = document.add_paragraph(style="List Bullet")
                run = paragraph.add_run(line[2:])
                run.font.size = Pt(11)
            elif line.strip():
                document.add_paragraph(line)

    document.save(output_path)
    return {"output_path": str(output_path), "status": "completed"}
